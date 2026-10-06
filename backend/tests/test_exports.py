from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.exports import build_export_workbook, process_export_job, process_export_jobs
from app.models import ExportJob


def _create_employee(client: TestClient, code: str, name: str, country: str = "IN", job_title_id: int = 1):
    response = client.post(
        "/api/v1/employees",
        json={
            "employee_code": code,
            "full_name": name,
            "job_title_id": job_title_id,
            "country_code": country,
            "salary": "1500000.00" if country == "IN" else "90000.00",
            "department": "Engineering",
            "employment_status": "active",
            "hired_at": "2024-01-01",
        },
    )
    assert response.status_code == 201


def test_queue_export_snapshots_current_filters(client: TestClient, db: Session):
    response = client.post(
        "/api/v1/exports",
        json={
            "recipient_email": "reviewer@example.com",
            "search": " Aryan ",
            "country_code": "in",
            "job_title_id": 1,
            "sort_by": "full_name",
            "sort_dir": "desc",
        },
    )
    assert response.status_code == 202
    body = response.json()
    assert body["status"] == "queued"

    job = db.get(ExportJob, body["id"])
    assert job is not None
    assert job.search == "Aryan"
    assert job.country_code == "IN"
    assert job.job_title_id == 1
    assert job.sort_dir == "desc"


def test_workbook_contains_all_matching_filtered_rows(client: TestClient, db: Session):
    _create_employee(client, "EMP-1001", "Asha Rao", "IN", 1)
    _create_employee(client, "EMP-1002", "Amit Rao", "IN", 1)
    _create_employee(client, "EMP-1003", "Alex Smith", "US", 1)

    queued = client.post(
        "/api/v1/exports",
        json={
            "recipient_email": "reviewer@example.com",
            "search": "Rao",
            "country_code": "IN",
            "sort_by": "full_name",
            "sort_dir": "asc",
        },
    ).json()
    job = db.get(ExportJob, queued["id"])
    assert job is not None

    workbook_bytes, row_count = build_export_workbook(db, job)
    assert row_count == 2
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True)
    rows = list(workbook["Employees"].iter_rows(values_only=True))
    assert rows[0][0] == "Employee Code"
    assert [row[1] for row in rows[1:]] == ["Amit Rao", "Asha Rao"]


def test_worker_marks_export_sent(client: TestClient, db: Session, monkeypatch):
    _create_employee(client, "EMP-2001", "Nina Patel")
    queued = client.post(
        "/api/v1/exports",
        json={"recipient_email": "reviewer@example.com", "country_code": "IN"},
    ).json()

    sent = {}

    def fake_send_export_email(**kwargs):
        sent.update(kwargs)

    monkeypatch.setattr("app.exports.send_export_email", fake_send_export_email)
    result = process_export_jobs(db)
    assert result["sent"] == 1

    job = db.get(ExportJob, queued["id"])
    assert job is not None
    assert job.status == "sent"
    assert job.row_count == 1
    assert sent["recipient_email"] == "reviewer@example.com"
    assert sent["xlsx_bytes"]


def test_worker_retries_delivery_failures(client: TestClient, db: Session, monkeypatch):
    queued = client.post(
        "/api/v1/exports",
        json={"recipient_email": "reviewer@example.com"},
    ).json()

    def fail_send(**_kwargs):
        raise RuntimeError("mail provider unavailable")

    monkeypatch.setattr("app.exports.send_export_email", fail_send)
    result = process_export_jobs(db)
    assert result["retried"] == 1

    job = db.get(ExportJob, queued["id"])
    assert job is not None
    assert job.status == "queued"
    assert job.attempts == 1
    assert "mail provider unavailable" in (job.last_error or "")


def test_same_export_cannot_be_delivered_twice(client: TestClient, db: Session, monkeypatch):
    queued = client.post(
        "/api/v1/exports",
        json={"recipient_email": "reviewer@example.com"},
    ).json()
    job = db.get(ExportJob, queued["id"])
    assert job is not None

    sends = []
    monkeypatch.setattr("app.exports.send_export_email", lambda **kwargs: sends.append(kwargs))

    assert process_export_job(db, job) == "sent"
    assert process_export_job(db, job) == "skipped"
    assert len(sends) == 1
